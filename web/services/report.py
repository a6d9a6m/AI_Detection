"""
报告生成服务
支持生成PDF、Excel、Word格式的异常检测报告
"""

import os
import json
import sys
from datetime import datetime
from typing import Optional, Dict, Any, List
import traceback

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image
from reportlab.lib import colors
from reportlab.pdfgen import canvas

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BarChart, Reference, LineChart

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# Add paths
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'src'))

from .db import Database, Report, Model
from paths import REPORTS as REPORTS_DIR, report_path as get_report_path
import matplotlib
matplotlib.use('Agg')  # 使用非GUI后端，避免线程问题
import matplotlib.pyplot as plt
import numpy as np


class ReportGenerator:
    """报告生成器"""

    def __init__(self):
        self.db = Database()
        self.report_dir = REPORTS_DIR

    def generate_report(self, user_id: int, tenant_id: Optional[int],
                       report_config: Dict[str, Any]) -> Dict[str, Any]:
        """
        生成异常检测报告 (TN-F-031)

        Args:
            user_id: 用户ID
            tenant_id: 租户ID
            report_config: 报告配置
                {
                    'title': str,  # 报告标题
                    'report_type': str,  # 报告类型
                    'model_id': int,  # 模型ID
                    'template_id': int,  # 模板ID (可选)
                    'format': str,  # 输出格式 (pdf/excel/word)
                    'content': dict,  # 报告内容
                    'language': str,  # 语言 (zh_CN/en_US)
                }

        Returns:
            包含report_id和文件路径的字典
        """
        try:
            # 验证配置
            if 'title' not in report_config:
                return {'error': 'Report title required'}

            if 'format' not in report_config:
                report_config['format'] = 'pdf'

            # 获取模型信息
            model = None
            if 'model_id' in report_config:
                model = Model.get_by_id(report_config['model_id'])

            # 准备报告数据
            report_data = self._prepare_report_data(report_config, model)

            # 根据格式生成报告
            format_type = report_config['format'].lower()

            if format_type == 'pdf':
                file_path = self._generate_pdf_report(report_data)
            elif format_type == 'excel':
                file_path = self._generate_excel_report(report_data)
            elif format_type == 'word':
                file_path = self._generate_word_report(report_data)
            else:
                return {'error': f'Unsupported format: {format_type}'}

            # 创建报告记录
            report_id = Report.create(
                user_id=user_id,
                tenant_id=tenant_id,
                title=report_config['title'],
                report_type=report_config.get('report_type', 'anomaly_detection'),
                model_id=report_config.get('model_id'),
                template_id=report_config.get('template_id'),
                content=report_data,
                file_path=file_path
            )

            return {
                'report_id': report_id,
                'file_path': file_path,
                'format': format_type,
                'message': 'Report generated successfully'
            }

        except Exception as e:
            traceback.print_exc()
            return {'error': str(e)}

    def _prepare_report_data(self, config: Dict[str, Any],
                            model: Optional[Dict] = None) -> Dict[str, Any]:
        """准备报告数据"""
        data = {
            'title': config['title'],
            'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'language': config.get('language', 'zh_CN'),
            'sections': []
        }

        # 添加模型信息部分
        if model:
            data['sections'].append({
                'title': '模型信息' if data['language'] == 'zh_CN' else 'Model Information',
                'type': 'model_info',
                'content': {
                    'name': model.get('name', 'N/A'),
                    'version': model.get('version', 'N/A'),
                    'machine_type': model.get('machine_type', 'N/A'),
                    'model_type': model.get('model_type', 'N/A'),
                    'created_at': model.get('created_at', 'N/A')
                }
            })

            # 添加性能指标部分
            metrics = model.get('metrics', {})
            if metrics:
                data['sections'].append({
                    'title': '性能指标' if data['language'] == 'zh_CN' else 'Performance Metrics',
                    'type': 'metrics',
                    'content': metrics
                })

        # 添加自定义内容部分
        if 'content' in config:
            data['sections'].append({
                'title': '详细内容' if data['language'] == 'zh_CN' else 'Detailed Content',
                'type': 'custom',
                'content': config['content']
            })

        # 添加统计数据部分（mock数据）
        data['sections'].append({
            'title': '统计摘要' if data['language'] == 'zh_CN' else 'Statistical Summary',
            'type': 'statistics',
            'content': {
                'total_samples': 1000,
                'normal_samples': 850,
                'anomaly_samples': 150,
                'detection_accuracy': 94.5,
                'false_positive_rate': 3.2,
                'false_negative_rate': 2.3
            }
        })

        return data

    def _generate_pdf_report(self, data: Dict[str, Any]) -> str:
        """生成PDF报告"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"report_{timestamp}.pdf"
        filepath = os.path.join(self.report_dir, filename)

        doc = SimpleDocTemplate(filepath, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()

        # 标题
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1f77b4'),
            spaceAfter=30,
            alignment=1  # Center
        )
        story.append(Paragraph(data['title'], title_style))
        story.append(Spacer(1, 0.2*inch))

        # 生成时间
        story.append(Paragraph(f"Generated: {data['generated_at']}", styles['Normal']))
        story.append(Spacer(1, 0.3*inch))

        # 各个部分
        for section in data['sections']:
            # 部分标题
            story.append(Paragraph(section['title'], styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))

            # 部分内容
            content = section['content']

            if section['type'] == 'model_info':
                table_data = [[k, str(v)] for k, v in content.items()]
                t = Table(table_data, colWidths=[2*inch, 3*inch])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, -1), colors.beige),
                    ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                story.append(t)

            elif section['type'] == 'metrics':
                table_data = [[k, str(v)] for k, v in content.items()]
                t = Table(table_data, colWidths=[2*inch, 3*inch])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, -1), colors.lightblue),
                    ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                story.append(t)

            elif section['type'] == 'statistics':
                # 创建统计图表
                fig, ax = plt.subplots(figsize=(6, 4))
                categories = ['Normal', 'Anomaly']
                values = [content.get('normal_samples', 0), content.get('anomaly_samples', 0)]
                ax.bar(categories, values, color=['green', 'red'])
                ax.set_ylabel('Number of Samples')
                ax.set_title('Sample Distribution')

                chart_filename = f'chart_{timestamp}.png'
                chart_path = os.path.join(self.report_dir, chart_filename)
                fig.savefig(chart_path)
                plt.close(fig)

                # 添加图表到报告
                story.append(Image(chart_path, width=4*inch, height=3*inch))

                # 添加统计表格
                table_data = [[k.replace('_', ' ').title(), str(v)] for k, v in content.items()]
                t = Table(table_data, colWidths=[2*inch, 3*inch])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, -1), colors.lightgreen),
                    ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                story.append(t)

            else:
                # 自定义内容
                content_text = json.dumps(content, indent=2)
                story.append(Paragraph(content_text, styles['Code']))

            story.append(Spacer(1, 0.3*inch))

        # 生成PDF
        doc.build(story)

        return filepath

    def _generate_excel_report(self, data: Dict[str, Any]) -> str:
        """生成Excel报告"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"report_{timestamp}.xlsx"
        filepath = os.path.join(self.report_dir, filename)

        wb = Workbook()
        ws = wb.active
        ws.title = "Report"

        # 标题
        ws['A1'] = data['title']
        ws['A1'].font = Font(size=18, bold=True, color="1F77B4")
        ws['A1'].alignment = Alignment(horizontal='center')
        ws.merge_cells('A1:D1')

        # 生成时间
        ws['A2'] = f"Generated: {data['generated_at']}"
        ws['A2'].font = Font(size=10)
        ws.merge_cells('A2:D2')

        row = 4

        # 各个部分
        for section in data['sections']:
            # 部分标题
            ws[f'A{row}'] = section['title']
            ws[f'A{row}'].font = Font(size=14, bold=True)
            ws[f'A{row}'].fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
            row += 1

            # 部分内容
            content = section['content']

            if isinstance(content, dict):
                for key, value in content.items():
                    ws[f'A{row}'] = key
                    ws[f'B{row}'] = str(value)
                    row += 1
            else:
                ws[f'A{row}'] = str(content)
                row += 1

            row += 1

        # 调整列宽
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 30

        wb.save(filepath)

        return filepath

    def _generate_word_report(self, data: Dict[str, Any]) -> str:
        """生成Word报告"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"report_{timestamp}.docx"
        filepath = os.path.join(self.report_dir, filename)

        doc = Document()

        # 标题
        title = doc.add_heading(data['title'], 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # 生成时间
        doc.add_paragraph(f"Generated: {data['generated_at']}")
        doc.add_paragraph()

        # 各个部分
        for section in data['sections']:
            # 部分标题
            doc.add_heading(section['title'], 1)

            # 部分内容
            content = section['content']

            if isinstance(content, dict):
                # 创建表格
                table = doc.add_table(rows=len(content) + 1, cols=2)
                table.style = 'Light Grid Accent 1'

                # 表头
                header_cells = table.rows[0].cells
                header_cells[0].text = 'Property'
                header_cells[1].text = 'Value'

                # 内容行
                for i, (key, value) in enumerate(content.items(), 1):
                    row_cells = table.rows[i].cells
                    row_cells[0].text = str(key)
                    row_cells[1].text = str(value)
            else:
                doc.add_paragraph(str(content))

            doc.add_paragraph()

        doc.save(filepath)

        return filepath

    def batch_generate_reports(self, user_id: int, tenant_id: Optional[int],
                               reports_config: List[Dict[str, Any]]) -> Dict[str, Any]:
        """
        批量生成报告 (TN-F-040)

        Args:
            user_id: 用户ID
            tenant_id: 租户ID
            reports_config: 报告配置列表

        Returns:
            批量生成结果
        """
        results = []
        errors = []

        for config in reports_config:
            try:
                result = self.generate_report(user_id, tenant_id, config)
                if 'error' in result:
                    errors.append({'config': config, 'error': result['error']})
                else:
                    results.append(result)
            except Exception as e:
                errors.append({'config': config, 'error': str(e)})

        return {
            'results': results,
            'errors': errors,
            'success_count': len(results),
            'error_count': len(errors)
        }


class ReportTemplateManager:
    """报告模板管理器"""

    def __init__(self):
        self.db = Database()

    def create_template(self, user_id: int, tenant_id: Optional[int],
                       template_config: Dict[str, Any]) -> int:
        """
        创建报告模板 (TN-F-032)

        Args:
            user_id: 用户ID
            tenant_id: 租户ID
            template_config: 模板配置

        Returns:
            模板ID
        """
        conn = self.db.get_connection()
        cursor = conn.cursor()

        cursor.execute('''
            INSERT INTO report_templates
            (user_id, tenant_id, name, description, template_content,
             fields, chart_config, style_config, language)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (user_id, tenant_id, template_config['name'],
              template_config.get('description'),
              json.dumps(template_config.get('template_content', {})),
              json.dumps(template_config.get('fields', [])),
              json.dumps(template_config.get('chart_config', {})),
              json.dumps(template_config.get('style_config', {})),
              template_config.get('language', 'zh_CN')))

        conn.commit()
        template_id = cursor.lastrowid
        conn.close()

        return template_id

    def get_template(self, template_id: int) -> Optional[Dict]:
        """获取报告模板"""
        conn = self.db.get_connection()
        cursor = conn.cursor()

        cursor.execute('SELECT * FROM report_templates WHERE id = ?', (template_id,))
        row = cursor.fetchone()
        conn.close()

        if row:
            result = dict(row)
            result['template_content'] = json.loads(result.get('template_content', '{}'))
            result['fields'] = json.loads(result.get('fields', '[]'))
            result['chart_config'] = json.loads(result.get('chart_config', '{}'))
            result['style_config'] = json.loads(result.get('style_config', '{}'))
            return result

        return None

    def list_templates(self, user_id: Optional[int] = None,
                      tenant_id: Optional[int] = None) -> List[Dict]:
        """列出报告模板"""
        conn = self.db.get_connection()
        cursor = conn.cursor()

        query = 'SELECT * FROM report_templates WHERE 1=1'
        params = []

        if user_id:
            query += ' AND user_id = ?'
            params.append(user_id)

        if tenant_id:
            query += ' AND tenant_id = ?'
            params.append(tenant_id)

        query += ' ORDER BY created_at DESC'

        cursor.execute(query, params)
        rows = cursor.fetchall()
        conn.close()

        results = []
        for row in rows:
            result = dict(row)
            result['template_content'] = json.loads(result.get('template_content', '{}'))
            result['fields'] = json.loads(result.get('fields', '[]'))
            result['chart_config'] = json.loads(result.get('chart_config', '{}'))
            result['style_config'] = json.loads(result.get('style_config', '{}'))
            results.append(result)

        return results


# 全局实例
report_generator = ReportGenerator()
template_manager = ReportTemplateManager()
